import sys
import numpy as np
import pandas as pd
import seaborn as sns
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def parse_sample(xlsx_path, sheet_name):
    wb = load_workbook(filename=xlsx_path, read_only=True)

    if sheet_name not in wb.sheetnames:
        raise ValueError(f"No sheet with name '{sheet_name}'")

    sheet = wb[sheet_name]

    count_tr = 0
    count_br = 0
    count_plus_tr = 0
    count_plus_br = 0
    count_trust_bam = 0
    count_trust_fastq = 0
    count_mixcr = 0

    for row in sheet.iter_rows(min_row=1, max_row=6):
        v_gene = row[1].value
        d_gene = row[2].value
        j_gene = row[3].value
        igv_check = row[4].value
        trust_bam = row[5].value
        trust_fastq = row[6].value
        mixcr = row[7].value
        if v_gene and isinstance(v_gene, str): 
            if v_gene.strip().startswith('TR'):
                count_tr += 1
                if igv_check.strip()[0] == "+":
                    count_plus_tr += 1
            if v_gene.strip().startswith('IG'):
                count_br += 1
                if igv_check.strip()[0] == "+":
                    count_plus_br += 1
        if d_gene and isinstance(d_gene, str):
            if d_gene.strip().startswith('TR'):
                count_tr += 1
                if igv_check.strip()[1] == "+":
                    count_plus_tr += 1
            if d_gene.strip().startswith('IG'):
                count_br += 1
                if igv_check.strip()[1] == "+":
                    count_plus_br += 1
        if j_gene and isinstance(j_gene, str):
            if j_gene.strip().startswith('TR'):
                count_tr += 1
                if igv_check.strip()[2] == "+":
                    count_plus_tr += 1
            if j_gene.strip().startswith('IG'):
                count_br += 1
                if igv_check.strip()[2] == "+":
                    count_plus_br += 1

        if trust_bam and trust_bam == "+":
            count_trust_bam += 1
        if trust_fastq and trust_fastq == "+":
            count_trust_fastq += 1
        if mixcr and mixcr == "+":
            count_mixcr += 1

    return count_tr, count_br, count_plus_tr, count_plus_br, count_trust_bam, count_trust_fastq, count_mixcr


def create_stat_table(samples_list, samples_clones, cov_file, outfile):

    cov_df = pd.read_csv(cov_file, sep="\t", names=['sample', 'coverage'], dtype={'sample':'string'})
    cov_df = cov_df.set_index('sample')

    with open(samples_list, "r") as f1:
        reader = f1.read()

    samples = reader.strip().split("\n")

    res_df = pd.DataFrame(columns=['sample', 'BCR_mixcr', 'TCR_mixcr', 'BCR_igv', 'TCR_igv', 'TRUST4_BAM', 'TRUST4_FASTQ', 'MiXCR', 'Precision_BCR', 'Precision_TCR', 'Precision_TRUST4_BAM', 'Precision_TRUST4_FASTQ', 'Precision_MiXCR', 'Median Coverage'])
    for sample in samples:
        try:
            tr,br,plus_tr,plus_br,plus_trust_bam,plus_trust_fastq,plus_mixcr = parse_sample(samples_clones, sample)
        except Exception as e:
            print(e)
        else:
            prec_br = pd.NA if br == 0 else plus_br / br
            prec_tr = pd.NA if tr == 0 else plus_tr / tr
            prec_trust_bam = plus_trust_bam / 5
            prec_trust_fastq = plus_trust_fastq / 5
            prec_mixcr = plus_mixcr / 5
            #res_df.loc[len(res_df)] = [sample, br, tr, plus_br, plus_tr, plus_trust_bam,plus_trust_fastq,plus_mixcr,prec_br, prec_tr,prec_trust_bam,prec_trust_fastq,prec_mixcr, cov_df.loc[sample, 'coverage']]
            res_df.loc[len(res_df)] = [sample, br, tr, plus_br, plus_tr, plus_trust_bam,plus_trust_fastq,plus_mixcr,prec_br, prec_tr,prec_trust_bam,prec_trust_fastq,prec_mixcr]

    res_df.to_csv(outfile,sep="\t", index=False)

def create_heatmap(infile, outfile):
    df = pd.read_csv(infile, sep="\t", dtype={'sample':'string'})
    df = df.set_index('sample')
    data = df.iloc[:, 7:12]

    mask = data.isna()

    plt.figure(figsize=(8, 10))
    ax = sns.heatmap(data,mask=mask, annot=True, fmt='.2f', cmap="Reds", linewidths=0.5)
    ax.axvline(x=2, color='white', linewidth=10)
    ax.text(1, -0.1, 'MiXCR pipeline relative to IGV', transform=ax.get_xaxis_transform(), 
        ha='center', va='bottom', fontsize=8, fontweight='bold')
    ax.text(3.5, -0.1, 'TRUST4/MiXCR relative to MiXCR pipeline', transform=ax.get_xaxis_transform(),
        ha='center', va='bottom', fontsize=8, fontweight='bold')
    plt.xlabel("Precision", labelpad=30, fontweight='bold')
    # remove y labels
    #plt.yticks(ticks=[], labels=[])

    plt.xticks(rotation=20, ticks=[0.5, 1.5, 2.5, 3.5, 4.5], labels=['BCR', 'TCR', 'TRUST4_BAM', 'TRUST4_FASTQ', 'MiXCR'])
    plt.tight_layout()
    plt.savefig(outfile, dpi=800)

def scatter_plot(infile, outfile):
    df = pd.read_csv(infile, sep="\t", dtype={'sample':'string'})
    df_long = pd.concat([
        df[['Median Coverage','Precision_TRUST4_BAM']].assign(Program='TRUST4_BAM').rename(columns={'Precision_TRUST4_BAM': 'Precision'}),
        df[['Median Coverage', 'Precision_MiXCR']].assign(Program='MiXCR').rename(columns={'Precision_MiXCR': 'Precision'})
    ], ignore_index=True)

    plt.figure(figsize=(8, 6))
    sns.scatterplot(
      data=df_long,
      x='Median Coverage',
      y='Precision',
      hue='Program',
      palette={'MiXCR': '#A90606', 'TRUST4_BAM': '#004777'},
      s=60
    )

    sns.regplot(data=df, x='Median Coverage', y='Precision_TRUST4_BAM', scatter=False, color='#004777')
    sns.regplot(data=df, x='Median Coverage', y='Precision_MiXCR', scatter=False, color='#A90606')

    plt.xlabel('Median Coverage')
    plt.ylabel('Precision')
    plt.legend(title='Tool')
    plt.grid(True, alpha=0.3)
    plt.savefig(outfile, dpi=800)


# run
create_stat_table("samples_names.txt", "samples_clones.xlsx", "samples_coverage.csv", "mixcr_igv_statistics.csv")

# plot
create_heatmap("mixcr_igv_statistics.csv", "heatmap.png")

scatter_plot("mixcr_igv_statistics.csv", "scatterplot.png")
